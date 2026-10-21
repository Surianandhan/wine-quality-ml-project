"""
docxkit.py -- small helpers for writing Word documents in the SASTRA report and
paper formats with python-docx.

A sample document is opened as the template so its page set-up, styles,
numbering definitions and footers are reused exactly; its body is then emptied
and the new content is appended element by element. Equations are typeset by
pandoc (LaTeX -> Office Math), so they stay editable in Word.
"""
import copy
import re
import subprocess
import tempfile
import zipfile
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches
from lxml import etree

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
M = "http://schemas.openxmlformats.org/officeDocument/2006/math"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
WP = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
NSDECL = f'xmlns:w="{W}" xmlns:m="{M}" xmlns:r="{R}"'


def X(xml):
    """Parse a WordprocessingML fragment (one element) with the w/m/r prefixes."""
    return etree.fromstring(f"<root {NSDECL}>{xml}</root>")[0]


def q(tag):
    p, t = tag.split(":")
    return "{%s}%s" % ({"w": W, "m": M, "r": R}[p], t)


# --------------------------------------------------------------------------
# Runs and inline markup
# --------------------------------------------------------------------------
def rpr(b=False, i=False, sz=None, font=None, color=None, u=False, caps=False, va=None, noproof=False):
    parts = []
    if font:
        parts.append(f'<w:rFonts w:ascii="{font}" w:eastAsia="{font}" w:hAnsi="{font}" w:cs="{font}"/>')
    if b:
        parts.append("<w:b/>")
    if i:
        parts.append("<w:i/>")
    if caps:
        parts.append("<w:smallCaps/>")
    if noproof:
        parts.append("<w:noProof/>")
    if color:
        parts.append(f'<w:color w:val="{color}"/>')
    if sz:
        parts.append(f'<w:sz w:val="{sz}"/>')
    if u:
        parts.append('<w:u w:val="single"/>')
    if va:
        parts.append(f'<w:vertAlign w:val="{va}"/>')
    return f"<w:rPr>{''.join(parts)}</w:rPr>" if parts else ""


def run(text, **fmt):
    out = []
    for k, piece in enumerate(text.split("\t")):
        if k:
            out.append(f"<w:r>{rpr(**fmt)}<w:tab/></w:r>")
        if piece:
            out.append(f'<w:r>{rpr(**fmt)}<w:t xml:space="preserve">{escape(piece)}</w:t></w:r>')
    return "".join(out)


_TOKEN = re.compile(r"(\*\*.+?\*\*|\*[^*\s][^*]*?\*|<sup>.+?</sup>|<sub>.+?</sub>)")


def md(text, **base):
    """Runs from light markup: **bold**, *italic*, <sup>x</sup>, <sub>x</sub>."""
    out = []
    for tok in _TOKEN.split(text):
        if not tok:
            continue
        fmt = dict(base)
        if tok.startswith("**") and tok.endswith("**"):
            fmt["b"] = True
            out.append(md(tok[2:-2], **fmt))
            continue
        if tok.startswith("<sup>"):
            fmt["va"] = "superscript"
            tok = tok[5:-6]
        elif tok.startswith("<sub>"):
            fmt["va"] = "subscript"
            tok = tok[5:-6]
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
            fmt["i"] = True
            tok = tok[1:-1]
        out.append(run(tok, **fmt))
    return "".join(out)


def ppr(jc=None, before=None, after=None, line=None, keep_next=False, ind_left=None, ind_right=None,
        hanging=None, first=None, style=None, num=None, shd=None, tabs=None, page_break_before=False):
    parts = []
    if style:
        parts.append(f'<w:pStyle w:val="{style}"/>')
    if keep_next:
        parts.append("<w:keepNext/>")
    if page_break_before:
        parts.append("<w:pageBreakBefore/>")
    if num:
        parts.append(f'<w:numPr><w:ilvl w:val="0"/><w:numId w:val="{num}"/></w:numPr>')
    if shd:
        parts.append(f'<w:shd w:val="clear" w:color="auto" w:fill="{shd}"/>')
    if tabs:
        parts.append("<w:tabs>" + "".join(f'<w:tab w:val="{v}" w:pos="{p}"/>' for v, p in tabs) + "</w:tabs>")
    sp = []
    if before is not None:
        sp.append(f'w:before="{before}"')
    if after is not None:
        sp.append(f'w:after="{after}"')
    if line is not None:
        sp.append(f'w:line="{line}" w:lineRule="auto"')
    if sp:
        parts.append(f"<w:spacing {' '.join(sp)}/>")
    ind = []
    if ind_left is not None:
        ind.append(f'w:left="{ind_left}"')
    if ind_right is not None:
        ind.append(f'w:right="{ind_right}"')
    if hanging is not None:
        ind.append(f'w:hanging="{hanging}"')
    if first is not None:
        ind.append(f'w:firstLine="{first}"')
    if ind:
        parts.append(f"<w:ind {' '.join(ind)}/>")
    if jc:
        parts.append(f'<w:jc w:val="{jc}"/>')
    return f"<w:pPr>{''.join(parts)}</w:pPr>" if parts else ""


def pageref(bookmark, cached, **fmt):
    r = rpr(**fmt)
    r2 = rpr(noproof=True, **fmt)
    return (f'<w:r>{r}<w:fldChar w:fldCharType="begin"/></w:r>'
            f'<w:r>{r}<w:instrText xml:space="preserve"> PAGEREF {bookmark} \\h </w:instrText></w:r>'
            f'<w:r>{r}<w:fldChar w:fldCharType="separate"/></w:r>'
            f'<w:r>{r2}<w:t>{escape(str(cached))}</w:t></w:r>'
            f'<w:r>{r}<w:fldChar w:fldCharType="end"/></w:r>')


# --------------------------------------------------------------------------
# Equations
# --------------------------------------------------------------------------
def latex_to_omml(display=(), inline=()):
    """Typeset LaTeX with pandoc. Returns (display oMathPara list, inline oMath list)."""
    lines = [f"$${e}$$\n" for e in display] + [f"${e}$\n" for e in inline]
    with tempfile.TemporaryDirectory() as tmp:
        src, out = Path(tmp) / "eq.md", Path(tmp) / "eq.docx"
        src.write_text("\n".join(lines), encoding="utf-8")
        subprocess.run(["pandoc", str(src), "-o", str(out)], check=True)
        root = etree.fromstring(zipfile.ZipFile(out).read("word/document.xml"))
    paras = root.findall(f".//{{{W}}}body/{{{W}}}p")
    disp, inl = [], []
    for k, p in enumerate(paras):
        if k < len(display):
            mp = p.find(f".//{{{M}}}oMathPara")
            if mp is None:
                raise ValueError(f"pandoc did not typeset: {display[k]}")
            disp.append(mp)
        else:
            om = p.find(f".//{{{M}}}oMath")
            if om is None:
                raise ValueError(f"pandoc did not typeset: {inline[k - len(display)]}")
            inl.append(om)
    if len(disp) != len(display) or len(inl) != len(inline):
        raise ValueError("equation count mismatch")
    # pandoc writes two constructs the Office Math schema rejects; normalise them.
    for root_el in disp + inl:
        for rp in root_el.iter(f"{{{M}}}rPr"):
            if rp.find(f"{{{M}}}nor") is not None:          # m:nor and m:sty are alternatives
                for sty in rp.findall(f"{{{M}}}sty"):
                    rp.remove(sty)
        for mc in root_el.iter(f"{{{M}}}mcPr"):              # schema order: count, then mcJc
            count = mc.find(f"{{{M}}}count")
            if count is not None:
                mc.remove(count)
                mc.insert(0, count)
    for mp in disp:
        pr = mp.find(f"{{{M}}}oMathParaPr")
        if pr is None:
            pr = etree.SubElement(mp, f"{{{M}}}oMathParaPr")
            mp.insert(0, pr)
        for jc in pr.findall(f"{{{M}}}jc"):
            pr.remove(jc)
        jc = etree.SubElement(pr, f"{{{M}}}jc")
        jc.set(f"{{{M}}}val", "center")
    return disp, inl


# --------------------------------------------------------------------------
# Document
# --------------------------------------------------------------------------
class Doc:
    def __init__(self, template):
        self.doc = Document(template)
        body = self.doc.element.body
        self.body = body
        # Section properties of the template, reused for our own section breaks.
        self.para_sects = [copy.deepcopy(s) for s in body.findall(f"./{{{W}}}p/{{{W}}}pPr/{{{W}}}sectPr")]
        for el in list(body):
            if el.tag != q("w:sectPr"):
                body.remove(el)
        part = self.doc.part
        for rid, rel in list(part.rels.items()):
            if rel.reltype in (RT.IMAGE, RT.HYPERLINK) or rel.reltype.endswith("/customXml"):
                part.drop_rel(rid)
        self._bm = 1000
        self._pic = 1000
        self.numbering = part.numbering_part.element
        self._links = {}

    # -- low level ----------------------------------------------------------
    def add(self, el):
        self.body.insert(len(self.body) - 1, el)
        return el

    def p(self, content="", pp=""):
        return self.add(X(f"<w:p>{pp}{content}</w:p>"))

    def bookmark(self, name, content):
        self._bm += 1
        return (f'<w:bookmarkStart w:id="{self._bm}" w:name="{name}"/>{content}'
                f'<w:bookmarkEnd w:id="{self._bm}"/>')

    def page_break(self):
        self.p('<w:r><w:br w:type="page"/></w:r>')

    def section_break(self, sect):
        self.add(X(f"<w:p><w:pPr>{etree.tostring(sect, encoding='unicode')}</w:pPr></w:p>"))

    def link(self, url, text, **fmt):
        if url not in self._links:
            self._links[url] = self.doc.part.relate_to(url, RT.HYPERLINK, is_external=True)
        return f'<w:hyperlink r:id="{self._links[url]}">{run(text, color="0563C1", u=True, **fmt)}</w:hyperlink>'

    def new_list(self, abstract_id):
        """A fresh w:num on an existing abstract list, so numbering restarts at 1."""
        ids = [int(n.get(q("w:numId"))) for n in self.numbering.findall(q("w:num"))]
        nid = max(ids) + 1
        self.numbering.append(X(
            f'<w:num w:numId="{nid}"><w:abstractNumId w:val="{abstract_id}"/>'
            f'<w:lvlOverride w:ilvl="0"><w:startOverride w:val="1"/></w:lvlOverride></w:num>'))
        return nid

    def picture(self, path, width_in, pp):
        para = self.doc.add_paragraph()            # python-docx appends before sectPr
        para.add_run().add_picture(str(path), width=Inches(width_in))
        el = para._p
        self._pic += 1
        for dp in el.iter(f"{{{WP}}}docPr"):
            dp.set("id", str(self._pic))
        r = el.find(q("w:r"))
        r.insert(0, X(rpr(noproof=True)))
        if pp:
            el.insert(0, X(pp))
        return el

    # -- tables -------------------------------------------------------------
    @staticmethod
    def cell(width, content, fill=None, pp_xml="", valign="center", span=1):
        tc = f'<w:tcW w:w="{width}" w:type="dxa"/>'
        if span > 1:
            tc += f'<w:gridSpan w:val="{span}"/>'
        if fill:
            tc += f'<w:shd w:val="clear" w:color="auto" w:fill="{fill}"/>'
        tc += f'<w:vAlign w:val="{valign}"/>'
        return f"<w:tc><w:tcPr>{tc}</w:tcPr><w:p>{pp_xml}{content}</w:p></w:tc>"

    def grid_table(self, header, rows, widths, sz=20, keep_next=None, align=None, bold_rows=(),
                   head_fill="E7E6E6", repeat_header=True):
        """Bordered table in the samples' style: grey bold header, fixed widths, rows that never split."""
        if keep_next is None:
            keep_next = len(rows) <= 14
        align = align or ["left"] * len(widths)
        total = sum(widths)
        kn = "<w:keepNext/>" if keep_next else ""

        def cpp(a):
            jc = f'<w:jc w:val="{a}"/>' if a != "left" else ""
            return f'<w:pPr>{kn}<w:spacing w:before="20" w:after="20" w:line="240" w:lineRule="auto"/>{jc}</w:pPr>'

        def row(cells, head=False, bold=False):
            tr = "<w:trPr><w:cantSplit/>" + ("<w:tblHeader/>" if head and repeat_header else "") + '<w:jc w:val="center"/></w:trPr>'
            tcs = []
            for k, (w, c) in enumerate(zip(widths, cells)):
                a = "center" if head else align[k]
                content = c if isinstance(c, str) and c.startswith("<w:") else md(str(c), sz=sz, b=head or bold)
                tcs.append(self.cell(w, content, head_fill if head else None, cpp(a)))
            return f"<w:tr>{tr}{''.join(tcs)}</w:tr>"

        grid = "".join(f'<w:gridCol w:w="{w}"/>' for w in widths)
        xml = ('<w:tbl><w:tblPr><w:tblStyle w:val="TableGrid"/>'
               f'<w:tblW w:w="{total}" w:type="dxa"/><w:jc w:val="center"/><w:tblLayout w:type="fixed"/>'
               '<w:tblLook w:val="04A0" w:firstRow="1" w:lastRow="0" w:firstColumn="1" w:lastColumn="0" w:noHBand="0" w:noVBand="1"/>'
               f"</w:tblPr><w:tblGrid>{grid}</w:tblGrid>")
        if header:
            xml += row(header, head=True)
        for k, r in enumerate(rows):
            xml += row(r, bold=k in bold_rows)
        xml += "</w:tbl>"
        return self.add(X(xml))

    def eq_table(self, maths, labels, widths=(1080, 7589, 1080), pp_xml=None):
        """Borderless 3-column table: empty | centred equation | (number)."""
        pp_xml = pp_xml or '<w:pPr><w:spacing w:before="60" w:after="60" w:line="240" w:lineRule="auto"/></w:pPr>'
        pr_right = '<w:pPr><w:spacing w:before="60" w:after="60" w:line="240" w:lineRule="auto"/><w:jc w:val="right"/></w:pPr>'
        grid = "".join(f'<w:gridCol w:w="{w}"/>' for w in widths)
        tbl = X('<w:tbl><w:tblPr><w:tblW w:w="0" w:type="auto"/><w:jc w:val="center"/>'
                '<w:tblBorders><w:top w:val="nil"/><w:left w:val="nil"/><w:bottom w:val="nil"/><w:right w:val="nil"/>'
                '<w:insideH w:val="nil"/><w:insideV w:val="nil"/></w:tblBorders><w:tblLayout w:type="fixed"/>'
                '<w:tblLook w:val="04A0" w:firstRow="1" w:lastRow="0" w:firstColumn="1" w:lastColumn="0" w:noHBand="0" w:noVBand="1"/>'
                f"</w:tblPr><w:tblGrid>{grid}</w:tblGrid></w:tbl>")
        keep = lambda x: x.replace("<w:pPr>", "<w:pPr><w:keepNext/>", 1)
        for k, (math, label) in enumerate(zip(maths, labels)):
            last = k == len(maths) - 1
            pl, pr = (pp_xml, pr_right) if last else (keep(pp_xml), keep(pr_right))
            tr = X('<w:tr><w:trPr><w:cantSplit/><w:jc w:val="center"/></w:trPr>'
                   + self.cell(widths[0], "", None, pl)
                   + self.cell(widths[1], "", None, pl)
                   + self.cell(widths[2], run(f"({label})"), None, pr) + "</w:tr>")
            tr.findall(q("w:tc"))[1].find(q("w:p")).append(copy.deepcopy(math))
            tbl.append(tr)
        return self.add(tbl)

    # -- output -------------------------------------------------------------
    def save(self, path, title, author, subject="", keywords=""):
        cp = self.doc.core_properties
        cp.title, cp.author, cp.last_modified_by = title, author, author
        cp.subject, cp.keywords, cp.comments, cp.category = subject, keywords, "", ""
        cp.revision = 1
        self.doc.save(str(path))
